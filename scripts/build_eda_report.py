"""Construye el informe académico del EDA en formato Word."""

from __future__ import annotations

from datetime import date
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "data" / "raw"
FIGURES = ROOT / "reports" / "figures"
TABLES = ROOT / "reports" / "tables"
OUTPUT = ROOT / "reports" / "Informe_EDA_Pulso_TransMi_Nao.docx"

NAVY = "17365D"
LIGHT_BLUE = "DCE6F1"
PALE_BLUE = "F2F6FA"
GRAY = "D9D9D9"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = GRAY) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=90, start=90, bottom=90, end=90) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Página ")
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.text = str(header)
        set_cell_shading(cell, NAVY)
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for row_idx, row in enumerate(rows):
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = str(value)
            cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_idx % 2 == 1:
                set_cell_shading(cells[idx], PALE_BLUE)
            for run in cells[idx].paragraphs[0].runs:
                run.font.size = Pt(8.5)
            cells[idx].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT if idx == 0 else WD_ALIGN_PARAGRAPH.CENTER
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_border(cell)
            set_cell_margins(cell)
            if widths:
                cell.width = Inches(widths[idx])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_figure(doc: Document, filename: str, caption: str, width: float = 6.35) -> None:
    image_path = FIGURES / filename
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_together = True
    paragraph.add_run().add_picture(str(image_path), width=Inches(width))
    cap = doc.add_paragraph(style="Caption")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.keep_with_next = False
    cap.add_run(caption)


def add_bullet(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.add_run(text)


def add_numbered(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="List Number")
    paragraph.add_run(text)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    for name, size in (("Title", 24), ("Heading 1", 16), ("Heading 2", 12.5)):
        style = styles[name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = True
    styles["Heading 1"].paragraph_format.space_before = Pt(12)
    styles["Heading 1"].paragraph_format.space_after = Pt(7)
    styles["Heading 1"].paragraph_format.keep_with_next = True
    styles["Heading 2"].paragraph_format.space_before = Pt(9)
    styles["Heading 2"].paragraph_format.space_after = Pt(5)
    styles["Heading 2"].paragraph_format.keep_with_next = True
    styles["Caption"].font.name = "Aptos"
    styles["Caption"].font.size = Pt(9)
    styles["Caption"].font.italic = True
    styles["Caption"].font.color.rgb = RGBColor(70, 70, 70)

    footer = section.footer
    add_page_number(footer.paragraphs[0])
    footer.paragraphs[0].runs[0].font.size = Pt(8)


def build_report() -> None:
    observations = pd.read_csv(RAW / "observations.csv", dtype={"station_id": "string"}, parse_dates=["observed_at"])
    context = pd.read_csv(RAW / "context.csv", parse_dates=["observed_at"])
    stations = pd.read_csv(RAW / "stations.csv", dtype={"station_id": "string"})
    data = observations.merge(stations, on="station_id").merge(context, on="observed_at")
    data["hour"] = data["observed_at"].dt.hour
    data["weekday_number"] = data["observed_at"].dt.dayofweek
    data["day_type"] = data["weekday_number"].ge(5).map({True: "Fin de semana", False: "Día laboral"})
    data["rain_level"] = pd.cut(
        data["rain_mm"],
        [-float("inf"), 0.1, 0.5, float("inf")],
        labels=["Muy baja", "Moderada", "Alta"],
    )
    data["event_active"] = data["event_intensity"].gt(0.1)

    station_summary = pd.read_csv(TABLES / "station_summary.csv")
    quality = pd.read_csv(TABLES / "quality_checks.csv")
    demand_stats = data["demand"].describe(percentiles=[0.25, 0.5, 0.75, 0.9, 0.95, 0.99])
    correlations = data[["demand", "rain_mm", "temperature_c", "event_intensity"]].corr()["demand"]
    weekday_means = data.groupby("weekday_number")["demand"].mean()
    day_names = ["Lunes", "Martes", "Miércoles", "Jueves", "Viernes", "Sábado", "Domingo"]
    type_means = data.groupby("day_type")["demand"].mean()
    rain_means = data.groupby("rain_level", observed=False)["demand"].mean()
    event_means = data.groupby("event_active")["demand"].mean()

    doc = Document()
    configure_document(doc)

    # Portada
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(88)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Análisis exploratorio de la demanda de Pulso TransMi")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_before = Pt(10)
    run = subtitle.add_run("Caracterización temporal espacial y contextual del conjunto inicial")
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(60, 60, 60)

    doc.add_paragraph().paragraph_format.space_before = Pt(80)
    metadata = [
        ("Proyecto", "Pulso TransMi"),
        ("Equipo", "Equipo Nao"),
        ("Programa", "Ciencia de Datos"),
        ("Institución", "Universidad Externado de Colombia"),
        ("Fecha", "18 de septiembre de 2026"),
    ]
    for label, value in metadata:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(f"{label}: ")
        r.bold = True
        p.add_run(value)

    doc.add_page_break()

    doc.add_heading("Resumen ejecutivo", level=1)
    doc.add_paragraph(
        "Este informe caracteriza el conjunto inicial del reto Pulso TransMi antes de construir modelos de pronóstico. "
        "El análisis cubre 51.840 observaciones de demanda de 12 estaciones, registradas cada 15 minutos durante "
        "45 días, entre el 26 de julio y el 8 de septiembre de 2026. También incorpora 4.320 registros temporales "
        "de lluvia, temperatura e intensidad de eventos."
    )
    doc.add_paragraph(
        "Los controles confirmaron que el conjunto está completo para el periodo declarado: no se detectaron "
        "duplicados, valores faltantes, demandas negativas, estaciones desconocidas ni intervalos ausentes. La "
        "demanda presenta una distribución asimétrica, diferencias amplias entre estaciones y dos picos diarios "
        "bien definidos. Las mayores medias aparecen a las 7:00 y entre las 17:00 y 18:00. Los días laborales "
        "superan claramente al fin de semana."
    )
    doc.add_paragraph(
        "Ricaurte - NQS registra la demanda media más alta, seguida por Banderas y Portal El Dorado. Las variables "
        "de contexto muestran asociaciones débiles o moderadas si se observan de forma aislada. Estas relaciones "
        "no deben interpretarse como efectos causales, porque también cambian con la hora, el día y la estación. "
        "El resultado principal es que el modelo deberá representar de manera explícita la estructura temporal, "
        "las diferencias por estación y los rezagos de demanda."
    )

    doc.add_heading("Contenido", level=2)
    for item in [
        "1 Objetivo y alcance",
        "2 Datos y metodología",
        "3 Calidad y consistencia",
        "4 Distribución de la demanda",
        "5 Comportamiento temporal",
        "6 Diferencias entre estaciones",
        "7 Variables de contexto",
        "8 Hallazgos para el modelado",
        "9 Limitaciones y conclusiones",
        "10 Referencias",
        "Anexo de tablas",
    ]:
        doc.add_paragraph(item)

    doc.add_page_break()
    doc.add_heading("1 Objetivo y alcance", level=1)
    doc.add_paragraph(
        "El objetivo del análisis exploratorio es comprender la estructura del conjunto de datos, evaluar su "
        "calidad e identificar patrones útiles para pronosticar la demanda de pasajeros por estación. El EDA busca "
        "responder cuatro preguntas: cómo se distribuye la demanda, cómo cambia en el tiempo, qué diferencias "
        "existen entre estaciones y qué relación inicial presenta con el contexto."
    )
    doc.add_paragraph(
        "La unidad de análisis es una observación de demanda para una estación en un intervalo de 15 minutos. El "
        "resultado no constituye todavía un modelo predictivo. Su función es orientar la construcción de variables, "
        "los baselines y la estrategia de validación temporal."
    )

    doc.add_heading("2 Datos y metodología", level=1)
    doc.add_heading("2.1 Fuentes y estructura", level=2)
    add_table(
        doc,
        ["Archivo", "Contenido", "Filas", "Clave"],
        [
            ["stations.csv", "Catálogo, corredor y coordenadas", "12", "station_id"],
            ["observations.csv", "Demanda observada cada 15 minutos", "51.840", "observed_at + station_id"],
            ["context.csv", "Lluvia, temperatura y eventos", "4.320", "observed_at"],
            ["metadata.json", "Cobertura, frecuencia y hashes", "No aplica", "dataset"],
        ],
        [1.2, 2.8, 0.8, 1.6],
    )
    doc.add_paragraph(
        "Las observaciones se integraron con el catálogo mediante station_id y con el contexto mediante observed_at. "
        "Los identificadores se conservaron como texto para proteger los ceros iniciales. Las fechas incluyen la "
        "zona horaria de Bogotá."
    )

    doc.add_heading("2.2 Preparación", level=2)
    for text in [
        "Conversión de observed_at a fecha y hora con zona horaria.",
        "Creación de hora, día de la semana y tipo de día.",
        "Clasificación de lluvia en muy baja, moderada y alta.",
        "Definición de evento activo cuando event_intensity es mayor que 0,1 para excluir colas numéricas casi nulas.",
        "Cálculo de resúmenes por estación, hora, día y contexto.",
    ]:
        add_numbered(doc, text)

    doc.add_heading("2.3 Criterios de interpretación", level=2)
    doc.add_paragraph(
        "Se utilizaron medias y medianas porque la distribución tiene cola derecha. Las gráficas temporales se "
        "agregaron por día u hora para hacer visibles los patrones. Las correlaciones son de Pearson y describen "
        "asociación lineal. No controlan simultáneamente por estación, hora o día."
    )

    doc.add_page_break()
    doc.add_heading("3 Calidad y consistencia", level=1)
    doc.add_paragraph(
        "Se ejecutaron trece verificaciones sobre volumen, unicidad, completitud, dominio y continuidad temporal. "
        "Todas obtuvieron resultado satisfactorio. Cada estación aporta 4.320 periodos y el contexto contiene los "
        "mismos 4.320 instantes esperados."
    )
    quality_rows = [[row.validacion, f"{int(row.resultado):,}".replace(",", "."), str(row.estado)] for row in quality.itertuples()]
    add_table(doc, ["Validación", "Resultado", "Estado"], quality_rows, [4.5, 1.0, 0.8])
    doc.add_paragraph(
        "La ausencia de fallas permite iniciar la etapa de modelado sin imputaciones ni eliminación de registros. "
        "Sin embargo, el pipeline deberá repetir estos controles cuando lleguen datos incrementales, porque la "
        "calidad del corte inicial no garantiza la de futuros lotes."
    )

    doc.add_page_break()
    doc.add_heading("4 Distribución de la demanda", level=1)
    stats_rows = [
        ["Observaciones", f"{int(demand_stats['count']):,}".replace(",", ".")],
        ["Media", f"{demand_stats['mean']:.1f}"],
        ["Mediana", f"{demand_stats['50%']:.1f}"],
        ["Desviación estándar", f"{demand_stats['std']:.1f}"],
        ["Percentil 75", f"{demand_stats['75%']:.1f}"],
        ["Percentil 95", f"{demand_stats['95%']:.1f}"],
        ["Máximo", f"{demand_stats['max']:.0f}"],
    ]
    add_table(doc, ["Estadístico", "Demanda"], stats_rows, [3.5, 2.0])
    doc.add_paragraph(
        "La media de 356,5 supera la mediana de 263, lo que confirma asimetría positiva. La mayor parte de los "
        "intervalos presenta demanda baja o media, mientras una proporción menor alcanza valores altos durante "
        "horas pico o en estaciones de mayor afluencia. El máximo de 2.284 no es por sí solo evidencia de error: "
        "aparece dentro de un patrón consistente de picos y pertenece a Ricaurte - NQS."
    )
    add_figure(
        doc,
        "00_distribucion_demanda.png",
        "Figura 1. Distribución de la demanda por intervalo. La visualización se limita al percentil 99 para evitar que la cola extrema comprima el histograma.",
    )

    doc.add_page_break()
    doc.add_heading("5 Comportamiento temporal", level=1)
    doc.add_heading("5.1 Evolución diaria", level=2)
    doc.add_paragraph(
        "La serie diaria repite un ciclo semanal estable. Los descensos más marcados coinciden con fines de semana, "
        "mientras los días laborales se mantienen en niveles superiores. El mayor total diario ocurre el 3 de "
        "agosto, con 468.376 registros de demanda acumulada entre estaciones; el menor corresponde al 26 de julio, "
        "con 332.887."
    )
    add_figure(doc, "01_demanda_diaria.png", "Figura 2. Demanda total diaria en las doce estaciones.")

    doc.add_heading("5.2 Patrón horario", level=2)
    doc.add_paragraph(
        "La demanda muestra dos máximos diarios. El primero se concentra entre las 6:00 y 8:00 y el segundo entre "
        "las 16:00 y 19:00. La hora con mayor promedio global es las 17:00, con 701,1 por estación e intervalo; "
        "le siguen las 18:00 y las 7:00. El nivel es menor durante la madrugada y desciende después de las 20:00."
    )
    add_figure(doc, "02_demanda_por_hora.png", "Figura 3. Demanda promedio por hora y tipo de día.")

    doc.add_page_break()
    doc.add_heading("5.3 Día de la semana", level=2)
    weekday_rows = [[day_names[i], f"{weekday_means.loc[i]:.1f}"] for i in range(7)]
    add_table(doc, ["Día", "Demanda promedio"], weekday_rows, [3.2, 2.2])
    doc.add_paragraph(
        f"Los días laborales promedian {type_means['Día laboral']:.1f}, frente a {type_means['Fin de semana']:.1f} "
        "durante el fin de semana. La diferencia es cercana al 28 %. Martes registra la media más alta, aunque las "
        "medias de lunes a jueves son muy próximas. Por tanto, el contraste laboral frente a fin de semana parece "
        "más relevante que pequeñas diferencias entre días laborales."
    )
    add_figure(doc, "04_mapa_calor_dia_hora.png", "Figura 4. Demanda promedio por día de la semana y hora.")

    doc.add_page_break()
    doc.add_heading("6 Diferencias entre estaciones", level=1)
    doc.add_paragraph(
        "La estación es una fuente central de heterogeneidad. Ricaurte - NQS promedia 683,7 por intervalo, más de "
        "tres veces el promedio de Portal Usme. Banderas y Portal El Dorado también se separan del resto. Estas "
        "diferencias justifican incorporar station_id como variable y evaluar el error por estación, no solamente "
        "una métrica global."
    )
    add_figure(doc, "03_demanda_por_estacion.png", "Figura 5. Demanda promedio por estación y corredor.")
    doc.add_paragraph(
        "La representación espacial permite comprobar la dispersión geográfica, pero no demuestra por sí sola un "
        "efecto de localización. El tamaño de los puntos corresponde a la demanda promedio y el color al corredor. "
        "Los niveles altos se concentran en estaciones de intercambio o alta conectividad, como Ricaurte y Banderas."
    )
    add_figure(doc, "07_mapa_estaciones.png", "Figura 6. Ubicación y demanda promedio de las estaciones.", width=5.2)

    doc.add_page_break()
    doc.add_heading("7 Variables de contexto", level=1)
    doc.add_heading("7.1 Lluvia y eventos", level=2)
    doc.add_paragraph(
        f"La demanda media es {rain_means['Muy baja']:.1f} con lluvia muy baja, {rain_means['Moderada']:.1f} con "
        f"lluvia moderada y {rain_means['Alta']:.1f} con lluvia alta. Para los eventos se utilizó un umbral de "
        f"intensidad superior a 0,1. Bajo esa definición, la media es {event_means[True]:.1f} durante eventos activos "
        f"y {event_means[False]:.1f} sin evento activo."
    )
    doc.add_paragraph(
        "Estas diferencias son descriptivas. La lluvia, los eventos y la demanda pueden coincidir con horas o días "
        "de mayor movimiento. Una evaluación predictiva deberá medir si estas variables mejoran el error fuera de "
        "muestra después de incluir la estructura temporal."
    )
    add_figure(doc, "05_demanda_lluvia_eventos.png", "Figura 7. Distribución de la demanda según nivel de lluvia y presencia de eventos activos.")

    doc.add_heading("7.2 Temperatura y correlaciones", level=2)
    doc.add_paragraph(
        f"La correlación lineal entre demanda y temperatura es {correlations['temperature_c']:.2f}; con intensidad "
        f"de eventos es {correlations['event_intensity']:.2f}; y con lluvia es {correlations['rain_mm']:.2f}. "
        "La temperatura presenta la asociación más visible, pero continúa siendo débil. Parte de esta relación "
        "puede reflejar el ciclo diario, ya que temperatura y demanda cambian según la hora."
    )
    add_figure(doc, "06_demanda_temperatura.png", "Figura 8. Relación entre temperatura observada y demanda.")
    add_figure(doc, "08_correlaciones.png", "Figura 9. Matriz de correlaciones lineales entre demanda y variables de contexto.", width=5.4)

    doc.add_page_break()
    doc.add_heading("8 Hallazgos para el modelado", level=1)
    doc.add_paragraph(
        "El EDA indica que un modelo útil deberá combinar patrones regulares con información reciente. Las "
        "siguientes decisiones se desprenden directamente de los resultados:"
    )
    for text in [
        "Incluir estación y corredor para representar niveles base diferentes.",
        "Codificar hora, día de la semana y tipo de día; para variables cíclicas conviene usar seno y coseno.",
        "Construir rezagos de 15 minutos, 1 hora, 1 día y 1 semana, además de medias móviles.",
        "Probar lluvia, temperatura e intensidad de eventos como variables externas y conservarlas solo si mejoran la validación temporal.",
        "Evaluar por estación y horizonte para evitar que un buen promedio oculte errores concentrados.",
        "Comparar contra persistencia, mismo intervalo del día anterior y mismo intervalo de la semana anterior.",
    ]:
        add_bullet(doc, text)

    doc.add_heading("8.1 Validación recomendada", level=2)
    doc.add_paragraph(
        "No debe aplicarse una partición aleatoria porque mezclaría pasado y futuro. Como primer corte se pueden usar "
        "los primeros 38 días para entrenamiento y los últimos 7 para validación. Después conviene realizar "
        "backtesting con varios cortes temporales. La métrica principal del reto es WAPE por estación y su promedio."
    )

    doc.add_heading("8.2 Riesgos que debe vigilar el pipeline", level=2)
    add_bullet(doc, "Cambios en el nivel de demanda por estación.")
    add_bullet(doc, "Pérdida de intervalos o llegada tardía de datos.")
    add_bullet(doc, "Nuevas categorías, estaciones o identificadores mal tipados.")
    add_bullet(doc, "Deterioro del desempeño reciente frente al acumulado.")

    doc.add_heading("9 Limitaciones y conclusiones", level=1)
    doc.add_heading("9.1 Limitaciones", level=2)
    doc.add_paragraph(
        "El conjunto cubre 45 días, periodo suficiente para observar varias repeticiones semanales, pero limitado "
        "para estudiar estacionalidad mensual, festivos o cambios de largo plazo. La demanda, el clima y los eventos "
        "son sintéticos. Por esta razón, los hallazgos describen el comportamiento del reto y no deben generalizarse "
        "directamente a la operación real de TransMilenio."
    )
    doc.add_paragraph(
        "El análisis es principalmente descriptivo. Las medias por lluvia o evento no controlan factores de "
        "confusión. La matriz de correlaciones solo captura relaciones lineales y no reemplaza la comparación de "
        "modelos en validación temporal."
    )

    doc.add_heading("9.2 Conclusiones", level=2)
    doc.add_paragraph(
        "El conjunto inicial es consistente y puede utilizarse para la fase de modelado. La demanda depende con "
        "fuerza de la estación, la hora y el tipo de día. Los dos picos diarios, la caída del fin de semana y las "
        "diferencias de nivel entre estaciones son los patrones más sólidos del EDA. Las variables de contexto "
        "pueden aportar información adicional, pero su utilidad debe demostrarse fuera de muestra."
    )
    doc.add_paragraph(
        "El siguiente paso es implementar baselines diarios y semanales, construir variables de rezago y comparar "
        "modelos mediante cortes temporales. Esta secuencia mantiene trazabilidad entre lo observado en el EDA y las "
        "decisiones posteriores del sistema MLOps."
    )

    doc.add_heading("10 Referencias", level=1)
    doc.add_paragraph(
        "Universidad Externado de Colombia. Pulso TransMi SDK para estudiantes. "
        "https://github.com/uexternadojz/pulso-transmi-sdk"
    )
    doc.add_paragraph(
        "Pulso TransMi API. Documentación interactiva y conjunto inicial. "
        "https://pulso-transmi.72-60-245-2.sslip.io/docs"
    )
    doc.add_paragraph(
        "Fuente geográfica declarada en los metadatos: servicio de estaciones troncales de TransMilenio. "
        "https://gis.transmilenio.gov.co/arcgis/rest/services/Troncal/consulta_estaciones_troncales/FeatureServer/0"
    )

    doc.add_page_break()
    doc.add_heading("Anexo A Resumen por estación", level=1)
    station_rows = []
    for row in station_summary.itertuples():
        station_rows.append(
            [
                row.station_name,
                row.corridor,
                f"{row.mean_demand:.1f}",
                f"{row.median_demand:.1f}",
                f"{row.max_demand:.0f}",
            ]
        )
    add_table(
        doc,
        ["Estación", "Corredor", "Media", "Mediana", "Máximo"],
        station_rows,
        [2.7, 1.3, 0.75, 0.75, 0.75],
    )

    core = doc.core_properties
    core.title = "Análisis exploratorio de la demanda de Pulso TransMi"
    core.subject = "EDA del conjunto inicial de Pulso TransMi"
    core.author = "Equipo Nao"
    core.keywords = "Pulso TransMi, EDA, demanda, MLOps, ciencia de datos"

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_report()
